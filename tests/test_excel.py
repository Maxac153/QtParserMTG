import io
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from mtgprices.errors import ImportFileError
from mtgprices.excel import HEADERS, export_workbook, read_links
from mtgprices.models import Card, LinkInput, Site
from tests.conftest import GF_URL, SCG_URL


def card(**overrides):
    values = dict(
        id=1,
        site=Site.GOLDFISH,
        url=GF_URL,
        quantity=2,
        name="Loran",
        set_name="BRO",
        price_usd=Decimal("6.96"),
        updated_at=None,
    )
    values.update(overrides)
    return Card(**values)


def xlsx(rows, title="Sheet"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_export_has_sheet_per_site_and_prices():
    data = export_workbook(
        {Site.STARCITYGAMES: [], Site.GOLDFISH: [card(), card(id=2, url=GF_URL + "2", price_usd=None)]},
        Decimal("90"),
    )
    workbook = load_workbook(io.BytesIO(data))
    assert workbook.sheetnames == ["StarCityGames", "MTGGoldfish"]
    rows = list(workbook["MTGGoldfish"].iter_rows(values_only=True))
    assert rows[0] == HEADERS
    assert rows[1] == (2, "Loran", "BRO", 6.96, 630, GF_URL)
    assert rows[2][3:5] == (None, None)


def test_export_never_writes_formulas():
    data = export_workbook({Site.GOLDFISH: [card(name='=HYPERLINK("http://evil","x")', set_name="+1+1")]}, Decimal(60))
    sheet = load_workbook(io.BytesIO(data))["MTGGoldfish"]
    assert sheet["B2"].data_type == "s"
    assert sheet["B2"].value == '=HYPERLINK("http://evil","x")'
    assert sheet["C2"].data_type == "s"


def test_roundtrip_export_import():
    data = export_workbook(
        {Site.GOLDFISH: [card()], Site.STARCITYGAMES: [card(site=Site.STARCITYGAMES, url=SCG_URL, quantity=1)]},
        Decimal(60),
    )
    items, errors = read_links(data)
    assert items == [LinkInput(2, GF_URL), LinkInput(1, SCG_URL)]
    assert errors == []


def test_import_legacy_format_and_errors():
    data = xlsx(
        [
            ["Количество карт", "Название", "Название сета", "Цена в рублях", "Цена в долларах", "Ссылка"],
            [3, "a", "b", 1, 1, GF_URL],
            ["4", "a", "b", 1, 1, SCG_URL],
            [None, "a", "b", 1, 1, GF_URL + "x"],
            [1.5, "a", "b", 1, 1, GF_URL + "y"],
            [1, "a", "b", 1, 1, "http://127.0.0.1/"],
            ["note", "без ссылки"],
        ]
    )
    items, errors = read_links(data)
    assert items == [LinkInput(3, GF_URL), LinkInput(4, SCG_URL), LinkInput(1, GF_URL + "x")]
    assert len(errors) == 2
    assert "строка 5" in errors[0]


def test_import_garbage():
    with pytest.raises(ImportFileError):
        read_links(b"not an excel file")


def test_import_row_limit(monkeypatch):
    monkeypatch.setattr("mtgprices.excel.MAX_IMPORT_ROWS", 3)
    with pytest.raises(ImportFileError, match="больше 3"):
        read_links(xlsx([[1, GF_URL]] * 5))
