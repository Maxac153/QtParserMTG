import io
import time
from decimal import Decimal
from urllib.parse import parse_qs, urlsplit

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from mtgprices.config import Settings
from mtgprices.errors import FetchError
from mtgprices.models import ParsedCard, Site
from mtgprices.web.app import create_app
from tests.conftest import GF_URL, SCG_URL

SCRYFALL_URL = "https://scryfall.com/card/bro/12/loran-of-the-third-path"

BASE = "http://127.0.0.1:8000"


class FakeRates:
    def __init__(self):
        self.rate: Decimal | Exception = Decimal("83.62")

    async def __call__(self) -> Decimal:
        if isinstance(self.rate, Exception):
            raise self.rate
        return self.rate


def make_client(tmp_path, fetcher, **settings):
    rates = FakeRates()
    app = create_app(
        Settings(data_dir=tmp_path, scheduler_interval=None, **settings), fetcher=fetcher, rate_fetcher=rates
    )
    client = TestClient(app, base_url=BASE, follow_redirects=False)
    client.fetcher = fetcher
    client.rates = rates
    client.repo = app.state.repo
    return client


@pytest.fixture
def client(tmp_path, fetcher):
    with make_client(tmp_path, fetcher) as client:
        yield client


def query_of(response) -> dict[str, str]:
    assert response.status_code == 303, response.text
    return {key: values[0] for key, values in parse_qs(urlsplit(response.headers["location"]).query).items()}


def wait_job(client, job_id: str) -> dict:
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] != "running":
            return job
        time.sleep(0.01)
    raise AssertionError("задача не завершилась")


def add_existing(client, site=Site.GOLDFISH, url=GF_URL, price="6.96", name="Loran"):
    return client.repo.add(site, url, 2, ParsedCard(name, "BRO", Decimal(price)))


def test_empty_index(client):
    response = client.get("/")
    assert response.status_code == 200
    assert "На этой вкладке пока нет карт" in response.text
    assert response.headers["x-frame-options"] == "DENY"
    assert "script-src 'self'" in response.headers["content-security-policy"]


def test_add_cards_flow(client):
    response = client.post("/cards/add", data={"links": f"3 {GF_URL}\n{SCG_URL}\nмусор\n"})
    params = query_of(response)
    assert params["site"] == "goldfish"
    job = wait_job(client, params["job"])
    assert (job["status"], job["total"], job["succeeded"], len(job["errors"])) == ("finished", 3, 2, 1)

    page = client.get("/", params={"site": "goldfish", "job": params["job"]}).text
    assert "Loran of the Third Path (Foil)" in page
    assert "The Brothers&#39; War" in page
    assert "готово" in page
    cards = client.get("/api/cards").json()
    assert {(c["site"], c["quantity"]) for c in cards} == {("goldfish", 3), ("starcitygames", 1)}
    gf = next(c for c in cards if c["site"] == "goldfish")
    assert (gf["price_usd"], gf["price_rub"]) == ("6.96", "420")


def test_add_duplicate_reports_error(client):
    add_existing(client)
    job = wait_job(client, query_of(client.post("/cards/add", data={"links": GF_URL}))["job"])
    assert job["succeeded"] == 0
    assert "уже есть" in job["errors"][0]
    assert client.fetcher.requested == []


def test_add_nothing(client):
    assert query_of(client.post("/cards/add", data={"links": "  \n"}))["notice"] == "nothing_to_add"


def test_html_is_escaped(client):
    add_existing(client, name="<script>alert(1)</script>")
    page = client.get("/?site=goldfish").text
    assert "<script>alert(1)</script>" not in page
    assert "&lt;script&gt;" in page


def test_search_and_totals(client):
    add_existing(client, name="Loran", price="6.96")
    add_existing(client, url=GF_URL + "2", name="Tithe", price="0.50")
    page = client.get("/", params={"site": "goldfish", "q": "tit"}).text
    assert ">Tithe<" in page and ">Loran<" not in page
    page = client.get("/", params={"site": "goldfish"}).text
    # 2 × 6.96 + 2 × 0.50 = 14.92 $; 2 × 7 × 60 + 2 × 1 × 60 = 960 ₽
    assert "14.92" in page and "960" in page


def test_rate(client):
    add_existing(client)
    assert query_of(client.post("/settings/rate", data={"rate": "92,5", "site": "goldfish"}))["notice"] == "rate_saved"
    assert client.get("/api/cards").json()[0]["price_rub"] == "647.50"
    assert query_of(client.post("/settings/rate", data={"rate": "-1"}))["notice"] == "rate_invalid"
    assert client.repo.get_rate() == Decimal("92.50")


def test_refresh_one_card(client):
    card = add_existing(client, price="1.00")
    assert query_of(client.post(f"/cards/{card.id}/refresh"))["notice"] == "card_refreshed"
    assert client.repo.get(card.id).price_usd == Decimal("6.96")

    client.fetcher.failing.add(GF_URL)
    assert query_of(client.post(f"/cards/{card.id}/refresh"))["notice"] == "card_refresh_failed"
    assert "404" in client.get("/?site=goldfish").text


def test_refresh_site_job(client):
    add_existing(client, price="1.00")
    add_existing(client, url=GF_URL + "2", price="1.00")
    add_existing(client, site=Site.STARCITYGAMES, url=SCG_URL, price="1.00")
    job = wait_job(client, query_of(client.post("/sites/goldfish/refresh"))["job"])
    assert (job["total"], job["succeeded"]) == (2, 2)
    prices = {c.site: c.price_usd for c in client.repo.list_cards()}
    assert prices == {Site.GOLDFISH: Decimal("6.96"), Site.STARCITYGAMES: Decimal("1.00")}


def test_refresh_empty_site(client):
    assert query_of(client.post("/sites/goldfish/refresh"))["notice"] == "nothing_to_refresh"


def test_quantity_and_delete(client):
    card = add_existing(client)
    assert query_of(client.post(f"/cards/{card.id}/quantity", data={"quantity": "5"}))["notice"] == "quantity_saved"
    assert client.repo.get(card.id).quantity == 5
    for bad in ("0", "abc", "1.5"):
        assert (
            query_of(client.post(f"/cards/{card.id}/quantity", data={"quantity": bad}))["notice"] == "quantity_invalid"
        )
    assert query_of(client.post(f"/cards/{card.id}/delete"))["notice"] == "card_deleted"
    assert client.repo.list_cards() == []
    assert client.post(f"/cards/{card.id}/delete").status_code == 404


def test_delete_all_on_site(client):
    add_existing(client)
    add_existing(client, site=Site.STARCITYGAMES, url=SCG_URL)
    client.post("/sites/goldfish/delete")
    assert [c.site for c in client.repo.list_cards()] == [Site.STARCITYGAMES]
    assert client.post("/sites/unknown/delete").status_code == 422


def test_export_and_import_roundtrip(client):
    add_existing(client)
    response = client.get("/export.xlsx")
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook["MTGGoldfish"]["F2"].value == GF_URL

    client.post("/sites/goldfish/delete")
    params = query_of(client.post("/cards/import", files={"file": ("cards.xlsx", response.content)}))
    job = wait_job(client, params["job"])
    assert job["succeeded"] == 1
    assert client.repo.list_cards()[0].quantity == 2


def test_import_rejects_garbage_and_large_files(client, tmp_path):
    assert query_of(client.post("/cards/import", files={"file": ("x.xlsx", b"garbage")}))["notice"] == "import_invalid"
    big = b"x" * (5 * 1024 * 1024 + 1)
    assert query_of(client.post("/cards/import", files={"file": ("x.xlsx", big)}))["notice"] == "import_invalid"
    workbook = Workbook()
    workbook.active.append(["нет ссылок"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    assert (
        query_of(client.post("/cards/import", files={"file": ("x.xlsx", buffer.getvalue())}))["notice"]
        == "nothing_to_add"
    )


def test_cancel_unknown_job(client):
    assert client.post("/jobs/nope/cancel").status_code == 404
    assert client.get("/api/jobs/nope").status_code == 404


def test_unknown_notice_is_ignored(client):
    assert "notice-" not in client.get("/?notice=<b>hi</b>").text


# --- защита локального сайта ---


@pytest.mark.parametrize(
    "headers",
    [
        {"Origin": "https://evil.example"},
        {"Origin": "null"},
        {"Referer": "https://evil.example/page"},
        {"Sec-Fetch-Site": "cross-site"},
        {"Sec-Fetch-Site": "same-site", "Origin": "http://127.0.0.1:8000"},
        {"Origin": "http://127.0.0.1:9999"},
    ],
)
def test_cross_site_post_is_blocked(client, headers):
    add_existing(client)
    response = client.post("/sites/goldfish/delete", headers=headers)
    assert response.status_code == 403
    assert len(client.repo.list_cards()) == 1


@pytest.mark.parametrize(
    "headers",
    [
        {"Origin": BASE, "Sec-Fetch-Site": "same-origin"},
        {"Referer": BASE + "/?site=goldfish"},
        {},
    ],
)
def test_same_origin_post_is_allowed(client, headers):
    response = client.post("/settings/rate", data={"rate": "70"}, headers=headers)
    assert response.status_code == 303


def test_foreign_host_header_is_rejected(client):
    # Защита от DNS rebinding: evil.example → 127.0.0.1
    response = client.get("/", headers={"Host": "evil.example"})
    assert response.status_code == 400
    assert client.get("/", headers={"Host": "localhost:8000"}).status_code == 200


def test_static_and_no_docs(client):
    assert client.get("/static/app.js").status_code == 200
    assert client.get("/docs").status_code == 404


# --- история цен, tix, Scryfall ---


def test_card_page_with_history_chart(client):
    card = add_existing(client, price="1.00")
    page = client.get(f"/cards/{card.id}")
    assert page.status_code == 200
    assert "График появится после второй проверки" in page.text
    client.post(f"/cards/{card.id}/refresh")
    page = client.get(f"/cards/{card.id}").text
    assert "<polyline" in page and "6.96" in page
    history = client.get(f"/api/cards/{card.id}/history").json()
    assert [h["price_usd"] for h in history] == ["1", "6.96"]
    assert client.get("/cards/999").status_code == 404


def test_price_change_and_tix_columns(client):
    card = add_existing(client, price="1.00")
    client.post(f"/cards/{card.id}/refresh")
    page = client.get("/?site=goldfish").text
    assert 'class="change up">+5.96' in page
    assert '<th class="num">tix</th>' in page and "1.15" in page
    # На вкладке без цен в tix колонки нет
    add_existing(client, site=Site.STARCITYGAMES, url=SCG_URL)
    assert '<th class="num">tix</th>' not in client.get("/?site=starcitygames").text


def test_add_scryfall_card(client):
    job = wait_job(client, query_of(client.post("/cards/add", data={"links": SCRYFALL_URL + "#foil"}))["job"])
    assert job["succeeded"] == 1
    assert client.fetcher.requested == ["https://api.scryfall.com/cards/bro/12"]
    card = client.get("/api/cards?site=scryfall").json()[0]
    assert (card["name"], card["price_usd"], card["price_tix"]) == ("Loran of the Third Path (Foil)", "6.33", "0.09")


# --- курс ЦБ и автообновление ---


def test_cbr_rate_button(client):
    assert query_of(client.post("/settings/rate/cbr", data={"site": "goldfish"}))["notice"] == "rate_saved"
    assert client.repo.get_rate() == Decimal("83.62")
    client.rates.rate = FetchError("нет сети")
    assert query_of(client.post("/settings/rate/cbr"))["notice"] == "rate_cbr_failed"
    assert client.repo.get_rate() == Decimal("83.62")


def test_auto_settings(client):
    assert (
        query_of(client.post("/settings/auto", data={"refresh_hours": "12", "auto_rate": "1"}))["notice"]
        == "auto_saved"
    )
    page = client.get("/").text
    assert "цены каждые 12 ч" in page and "курс ЦБ" in page
    assert query_of(client.post("/settings/auto", data={"refresh_hours": "12"}))["notice"] == "auto_saved"
    assert client.repo.get_setting("auto_rate") == "0"
    for bad in ("-1", "abc", "1000"):
        assert query_of(client.post("/settings/auto", data={"refresh_hours": bad}))["notice"] == "auto_invalid"


def test_scheduler_runs_refresh_inside_app(tmp_path, fetcher):
    with make_client(tmp_path, fetcher) as client:
        add_existing(client, price="1.00")
        client.post("/settings/auto", data={"refresh_hours": "1"})
        client.portal.call(client.app.state.scheduler.tick)
        deadline = time.monotonic() + 5
        while client.repo.list_cards()[0].price_usd != Decimal("6.96"):
            assert time.monotonic() < deadline
            time.sleep(0.01)


# --- задачи переживают перезапуск ---


def test_unfinished_job_resumes_on_startup(tmp_path, fetcher):
    with make_client(tmp_path, fetcher) as client:
        client.repo.create_job("left", "add", "Добавление карт (1)", [{"quantity": 2, "url": GF_URL}], [])
    with make_client(tmp_path, fetcher) as client:
        job = wait_job(client, "left")
        assert (job["status"], job["succeeded"]) == ("finished", 1)
        assert client.repo.list_cards()[0].quantity == 2


# --- пароль ---


def test_password_protects_everything(tmp_path, fetcher):
    with make_client(tmp_path, fetcher, password="s3cret") as client:
        response = client.get("/")
        assert response.status_code == 401
        assert response.headers["www-authenticate"].startswith("Basic")
        assert client.get("/static/app.js").status_code == 401
        assert client.get("/", auth=("any", "wrong")).status_code == 401
        assert client.get("/", headers={"Authorization": "Basic !!!"}).status_code == 401
        assert client.get("/", auth=("any", "s3cret")).status_code == 200
        assert client.post("/settings/rate", data={"rate": "70"}, auth=("u", "s3cret")).status_code == 303
