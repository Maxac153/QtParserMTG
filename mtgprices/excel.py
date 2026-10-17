import io
import zipfile
from collections.abc import Mapping, Sequence
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import TYPE_STRING
from openpyxl.utils import get_column_letter

from mtgprices.errors import ImportFileError, InvalidLinkError
from mtgprices.links import check_quantity, validate_link
from mtgprices.models import Card, LinkInput, Site
from mtgprices.pricing import ruble_price

HEADERS = ("Количество", "Название", "Сет", "Цена, $", "Цена, ₽", "Ссылка")
COLUMN_WIDTHS = (12, 40, 30, 10, 10, 60)
MAX_IMPORT_ROWS = 5000


def export_workbook(cards_by_site: Mapping[Site, Sequence[Card]], rate: Decimal) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    for site, cards in cards_by_site.items():
        sheet = workbook.create_sheet(site.label)
        sheet.append(HEADERS)
        for card in cards:
            price_rub = None if card.price_usd is None else float(ruble_price(card.price_usd, rate))
            price_usd = None if card.price_usd is None else float(card.price_usd)
            sheet.append([card.quantity, None, None, price_usd, price_rub, None])
            row = sheet.max_row
            # Текст с сайтов и ссылки пишем строго как строки: «=...» не должно стать формулой
            for column, value in ((2, card.name), (3, card.set_name), (6, card.url)):
                cell = sheet.cell(row=row, column=column)
                cell.value = value
                cell.data_type = TYPE_STRING
        for index, width in enumerate(COLUMN_WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def read_links(data: bytes) -> tuple[list[LinkInput], list[str]]:
    """Читает ссылки из xlsx: количество — первая колонка, ссылка — первая ячейка строки, похожая на URL.

    Подходит и для файлов, выгруженных самим приложением, и для старого формата.
    """
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ImportFileError("Файл не похож на книгу Excel (.xlsx)") from exc

    items: list[LinkInput] = []
    errors: list[str] = []
    rows_seen = 0
    try:
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                rows_seen += 1
                if rows_seen > MAX_IMPORT_ROWS:
                    raise ImportFileError(f"В файле больше {MAX_IMPORT_ROWS} строк")
                url = next((v.strip() for v in row if isinstance(v, str) and v.strip().startswith("http")), None)
                if url is None:
                    continue
                where = f"«{sheet.title}», строка {row_number}"
                try:
                    quantity = _quantity(row[0] if row else None)
                    validate_link(url)
                except InvalidLinkError as exc:
                    errors.append(f"{where}: {exc}")
                    continue
                items.append(LinkInput(quantity, url))
    finally:
        workbook.close()
    return items, errors


def _quantity(value: object) -> int:
    if value is None or value == "":
        return 1
    try:
        number = float(str(value).strip().replace(",", "."))
    except ValueError as exc:
        raise InvalidLinkError(f"Некорректное количество: {value!r}") from exc
    if not number.is_integer():
        raise InvalidLinkError(f"Некорректное количество: {value!r}")
    return check_quantity(int(number))
